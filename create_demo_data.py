"""
create_demo_data.py -- generate the SYNTHETIC DEMO DATA corpus.

Every file describes the same fictional project (Project Alpha) with facts
that deliberately overlap across modalities, so a single question can only
be answered well by combining them.

    project_report.pdf       background, 2024 results, budget       (PDF)
    project_notes.docx       meeting notes and risks                (DOCX)
    project_metrics.csv      quarterly metrics                      (CSV)
    project_dashboard.png    rendered dashboard with real text      (IMAGE)
    project_meeting.wav      spoken summary                         (AUDIO)

The WAV is produced with Windows SAPI (pyttsx3 or PowerShell
System.Speech), which is real synthesised speech that Whisper can actually
transcribe. If no speech engine is available the script says so plainly and
skips the audio file rather than writing silence and claiming success.
"""
from __future__ import annotations

import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

from config import settings

BANNER = "SYNTHETIC DEMO DATA - Project Alpha (fictional)"

FACTS = {
    "records": "1.2 million records",
    "accuracy": "99.4 percent",
    "budget": "87 percent",
    "milestones": "14",
    "team": "6 engineers",
    "deployment": "Q2 2024",
}

NARRATION = (
    "This is the Project Alpha quarterly review for 2024. "
    "The ingestion platform processed one point two million records this year. "
    "Our classification accuracy reached ninety nine point four percent, "
    "which is above the target we set in the first quarter. "
    "Budget utilisation stands at eighty seven percent. "
    "The natural language processing module was deployed in the second quarter of 2024. "
    "We completed fourteen milestones with a team of six engineers. "
    "The main outstanding risk is the latency of the batch pipeline under peak load."
)

REPORT_PAGES = [
    (
        "Project Alpha - Annual Report 2024",
        f"{BANNER}. "
        "Project Alpha is an internal data platform initiative launched in 2023 to "
        "consolidate records processing across departments. This report covers the "
        "2024 development cycle, including delivery milestones, accuracy results, "
        "budget utilisation and outstanding technical risks. The programme is "
        "governed by a steering committee that meets quarterly.",
    ),
    (
        "2024 Development Results",
        f"During 2024 the ingestion platform processed {FACTS['records']}, exceeding "
        f"the annual target of one million. Classification accuracy reached "
        f"{FACTS['accuracy']} on the held-out evaluation set, measured across four "
        f"quarterly checkpoints. The natural language processing module entered "
        f"production in {FACTS['deployment']}. A total of {FACTS['milestones']} "
        f"milestones were completed by a core team of {FACTS['team']}.",
    ),
    (
        "Budget and Resourcing",
        f"Budget utilisation for the 2024 financial year closed at {FACTS['budget']} "
        "of the allocated amount. Underspend was concentrated in infrastructure, "
        "where reserved capacity was reduced after the Q2 migration. Staffing "
        "remained stable throughout the year with no unplanned attrition.",
    ),
    (
        "Risks and Outlook",
        "The principal outstanding risk is batch pipeline latency under peak load, "
        "which exceeded the service objective during two separate load tests in "
        "Q4 2024. Mitigation work is scheduled for 2025. No security incidents "
        "were recorded during the reporting period.",
    ),
]


def make_pdf(path: Path) -> None:
    import pymupdf

    doc = pymupdf.open()
    for heading, body in REPORT_PAGES:
        page = doc.new_page()
        page.insert_text((60, 70), heading, fontsize=17, fontname="helv")
        page.insert_textbox(
            pymupdf.Rect(60, 100, 540, 700), body, fontsize=11, fontname="helv",
        )
        page.insert_text((60, 780), BANNER, fontsize=7, fontname="helv")
    doc.save(str(path))
    doc.close()


def make_docx(path: Path) -> None:
    import docx

    document = docx.Document()
    document.add_heading("Project Alpha - Internal Notes", level=1)
    document.add_paragraph(BANNER)

    document.add_heading("Q4 Review Meeting", level=2)
    document.add_paragraph(
        f"The team confirmed that {FACTS['records']} were processed during 2024 and "
        f"that accuracy held at {FACTS['accuracy']}. The steering committee accepted "
        f"the result without amendment."
    )
    document.add_paragraph(
        f"Budget utilisation was reported at {FACTS['budget']}. Finance asked for a "
        "written explanation of the infrastructure underspend before sign-off."
    )

    document.add_heading("Open Risks", level=2)
    document.add_paragraph(
        "Batch pipeline latency under peak load remains the top risk carried into "
        "2025. Two Q4 load tests breached the service objective."
    )

    document.add_heading("Milestone Summary", level=2)
    table = document.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text, header[1].text, header[2].text = "Quarter", "Milestones", "Status"
    for quarter, count, status in [
        ("Q1 2024", "3", "Complete"),
        ("Q2 2024", "4", "Complete - NLP module deployed"),
        ("Q3 2024", "4", "Complete"),
        ("Q4 2024", "3", "Complete"),
    ]:
        row = table.add_row().cells
        row[0].text, row[1].text, row[2].text = quarter, count, status

    document.save(str(path))


def make_csv(path: Path) -> None:
    import pandas as pd

    pd.DataFrame([
        {"quarter": "Q1 2024", "records_processed": 240000, "accuracy_pct": 97.1,
         "budget_used_pct": 19, "milestones": 3},
        {"quarter": "Q2 2024", "records_processed": 310000, "accuracy_pct": 98.3,
         "budget_used_pct": 44, "milestones": 4},
        {"quarter": "Q3 2024", "records_processed": 330000, "accuracy_pct": 99.0,
         "budget_used_pct": 67, "milestones": 4},
        {"quarter": "Q4 2024", "records_processed": 320000, "accuracy_pct": 99.4,
         "budget_used_pct": 87, "milestones": 3},
    ]).to_csv(path, index=False)


def make_dashboard(path: Path) -> None:
    """Render a dashboard with real, OCR-readable text and a bar chart."""
    from PIL import Image, ImageDraw, ImageFont

    width, height = 1100, 620
    image = Image.new("RGB", (width, height), "white")
    draw = ImageDraw.Draw(image)

    def font(size: int):
        for name in ("arial.ttf", "segoeui.ttf", "calibri.ttf"):
            try:
                return ImageFont.truetype(name, size)
            except OSError:
                continue
        return ImageFont.load_default()

    draw.rectangle([0, 0, width, 86], fill="#1f3a5f")
    draw.text((36, 22), "PROJECT ALPHA - 2024 DEVELOPMENT DASHBOARD",
              fill="white", font=font(26))

    metrics = [
        ("Records processed", "1.2 million"),
        ("Classification accuracy", "99.4%"),
        ("Budget utilisation", "87%"),
        ("Milestones completed", "14"),
        ("Team size", "6 engineers"),
        ("NLP module deployed", "Q2 2024"),
    ]
    for index, (label, value) in enumerate(metrics):
        x = 40 + (index % 2) * 300
        y = 130 + (index // 2) * 84
        draw.text((x, y), label, fill="#5a6472", font=font(16))
        draw.text((x, y + 26), value, fill="#16191f", font=font(28))

    draw.text((660, 130), "Records processed by quarter", fill="#5a6472", font=font(16))
    values = [240, 310, 330, 320]
    for index, value in enumerate(values):
        bar_height = int(value * 0.75)
        x0 = 670 + index * 92
        draw.rectangle([x0, 480 - bar_height, x0 + 62, 480], fill="#2e7d32")
        draw.text((x0 + 8, 490), f"Q{index + 1}", fill="#16191f", font=font(16))
        draw.text((x0 + 2, 460 - bar_height), f"{value}k", fill="#16191f", font=font(14))

    draw.text((36, height - 34), BANNER, fill="#8b94a0", font=font(14))
    image.save(path)


def make_audio(path: Path) -> tuple[bool, str]:
    """Real synthesised speech via Windows SAPI. Returns (ok, detail)."""
    try:
        import pyttsx3

        engine = pyttsx3.init()
        engine.setProperty("rate", 155)
        engine.save_to_file(NARRATION, str(path))
        engine.runAndWait()
        engine.stop()
        if path.exists() and path.stat().st_size > 8000:
            return True, "pyttsx3 (Windows SAPI)"
    except Exception as exc:
        print(f"  pyttsx3 unavailable ({exc}); trying PowerShell System.Speech")

    script = (
        "Add-Type -AssemblyName System.Speech; "
        "$s = New-Object System.Speech.Synthesis.SpeechSynthesizer; "
        "$s.Rate = 0; "
        f"$s.SetOutputToWaveFile('{path}'); "
        f"$s.Speak('{NARRATION}'); "
        "$s.Dispose()"
    )
    try:
        subprocess.run(
            ["powershell", "-NoProfile", "-NonInteractive", "-Command", script],
            capture_output=True, text=True, timeout=180, check=False,
        )
        if path.exists() and path.stat().st_size > 8000:
            return True, "PowerShell System.Speech (Windows SAPI)"
    except Exception as exc:
        return False, f"PowerShell speech failed: {exc}"

    return False, "No local speech engine produced audio."


def main() -> int:
    settings.ensure_dirs()
    documents = settings.documents_dir
    images = settings.images_dir
    audio = settings.audio_dir

    print(f"Creating {BANNER}\n")

    targets = [
        ("project_report.pdf", documents / "project_report.pdf", make_pdf),
        ("project_notes.docx", documents / "project_notes.docx", make_docx),
        ("project_metrics.csv", documents / "project_metrics.csv", make_csv),
        ("project_dashboard.png", images / "project_dashboard.png", make_dashboard),
    ]
    for name, path, builder in targets:
        builder(path)
        print(f"  OK   {name:24s} {path.stat().st_size:>8,} bytes")

    wav_path = audio / "project_meeting.wav"
    ok, detail = make_audio(wav_path)
    if ok:
        print(f"  OK   {'project_meeting.wav':24s} {wav_path.stat().st_size:>8,} bytes"
              f"  [{detail}]")
    else:
        print(f"  SKIP project_meeting.wav -- {detail}")
        print("\n  No synthetic audio was written. The audio pipeline is not")
        print("  demonstrated by this dataset. To test it, drop any real")
        print(f"  WAV/MP3/M4A file into {audio} and re-run ingestion.")

    print("\nDemo data ready.")
    return 0


if __name__ == "__main__":
    sys.exit(main())
