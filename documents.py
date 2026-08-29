"""
documents.py -- local extraction of PDF / DOC / DOCX / TXT / CSV into
ContentItem objects.

    PDF   PyMuPDF text per page, OCR fallback for scanned pages, page_number kept
    DOCX  python-docx paragraphs + tables, section kept (never a page number)
    DOC   LibreOffice headless -> DOCX -> the DOCX path above
    TXT   read and chunk
    CSV   pandas, one searchable row representation per row, row_index kept

Nothing here touches the network.
"""
from __future__ import annotations

import logging
import shutil
import subprocess
import tempfile
from pathlib import Path

from config import settings
from schemas import (
    ContentItem,
    ErrorCode,
    IngestionResult,
    Modality,
    Relationship,
    RelationshipType,
    Source,
    SourceLocation,
    SourceType,
    content_hash,
    derive_id,
)

logger = logging.getLogger(__name__)

EXTENSION_TO_TYPE = {
    ".pdf": SourceType.PDF,
    ".doc": SourceType.DOC,
    ".docx": SourceType.DOCX,
    ".txt": SourceType.TXT,
    ".md": SourceType.TXT,
    ".csv": SourceType.CSV,
}

# A page yielding fewer characters than this is treated as scanned and
# routed through OCR instead of being indexed as an empty page.
MIN_PAGE_CHARS_BEFORE_OCR = 80


# ---------------------------------------------------------------------------
# Chunking
# ---------------------------------------------------------------------------

def chunk_text(
    text: str,
    chunk_size: int | None = None,
    overlap: int | None = None,
) -> list[str]:
    """Split on a character window with overlap, preferring to break at a
    sentence or whitespace boundary so chunks stay readable as citations."""
    size = chunk_size or settings.chunk_size
    over = overlap if overlap is not None else settings.chunk_overlap
    text = " ".join(text.split())
    if not text:
        return []
    if len(text) <= size:
        return [text]

    chunks: list[str] = []
    start = 0
    while start < len(text):
        end = min(start + size, len(text))
        if end < len(text):
            window = text[start:end]
            # Prefer a sentence end, then any whitespace, in the last 30%.
            pivot = max(window.rfind(". "), window.rfind("! "), window.rfind("? "))
            if pivot < size * 0.7:
                pivot = window.rfind(" ")
            if pivot > size * 0.5:
                end = start + pivot + 1
        chunk = text[start:end].strip()
        if len(chunk) >= settings.min_chunk_chars or not chunks:
            chunks.append(chunk)
        if end >= len(text):
            break
        start = max(end - over, start + 1)
    return chunks


# ---------------------------------------------------------------------------
# Validation helpers
# ---------------------------------------------------------------------------

def resolve_source_type(path: Path) -> SourceType:
    return EXTENSION_TO_TYPE.get(path.suffix.lower(), SourceType.UNKNOWN)


def _validate(path_str: str) -> tuple[Path, SourceType]:
    path = Path(path_str).resolve()
    if not path.exists() or not path.is_file():
        raise FileNotFoundError(f"{ErrorCode.FILE_NOT_FOUND.value}: {path_str}")
    if path.stat().st_size > settings.max_file_bytes:
        raise ValueError(f"{ErrorCode.FILE_TOO_LARGE.value}: {path.name}")
    source_type = resolve_source_type(path)
    if source_type is SourceType.UNKNOWN:
        raise ValueError(f"{ErrorCode.UNSUPPORTED_FILE_TYPE.value}: {path.suffix}")
    return path, source_type


def _build_source(path: Path, source_type: SourceType) -> Source:
    data = path.read_bytes()
    return Source(
        source_id=content_hash(data),
        filename=path.name,
        source_type=source_type,
        file_path=str(path),
        mime_type=_mime_for(source_type),
        file_size=len(data),
    )


def _mime_for(source_type: SourceType) -> str:
    return {
        SourceType.PDF: "application/pdf",
        SourceType.DOC: "application/msword",
        SourceType.DOCX: (
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
        ),
        SourceType.TXT: "text/plain",
        SourceType.CSV: "text/csv",
    }.get(source_type, "application/octet-stream")


# ---------------------------------------------------------------------------
# Per-format extraction -> list of (text, SourceLocation, extra_metadata)
# ---------------------------------------------------------------------------

Extracted = tuple[str, SourceLocation, dict]


def _extract_pdf(path: Path, warnings: list[str]) -> list[Extracted]:
    import pymupdf

    out: list[Extracted] = []
    with pymupdf.open(path) as doc:
        for page_index, page in enumerate(doc, start=1):
            text = page.get_text("text").strip()
            ocr_used = False

            if len(text) < MIN_PAGE_CHARS_BEFORE_OCR:
                ocr_text = _ocr_pdf_page(page, warnings, page_index)
                if len(ocr_text) > len(text):
                    text, ocr_used = ocr_text, True

            if not text:
                continue
            for chunk in chunk_text(text):
                out.append((
                    chunk,
                    SourceLocation(page_number=page_index),
                    {"ocr_used": ocr_used},
                ))
    return out


def _ocr_pdf_page(page, warnings: list[str], page_index: int) -> str:
    """Rasterise one PDF page and OCR it. Returns '' on any failure -- a
    scanned page we cannot read is skipped, never fabricated."""
    try:
        import io

        import pytesseract
        from PIL import Image

        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
        pixmap = page.get_pixmap(dpi=200)
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        return pytesseract.image_to_string(image).strip()
    except Exception as exc:
        warnings.append(f"OCR failed on page {page_index}: {exc}")
        logger.warning("PDF OCR failed on page %s: %s", page_index, exc)
        return ""


def _extract_docx(path: Path, warnings: list[str]) -> list[Extracted]:
    """python-docx cannot know page boundaries -- pagination is computed by
    the renderer, not stored in the file -- so DOCX items carry a section
    heading and never a page_number."""
    import docx

    document = docx.Document(str(path))
    out: list[Extracted] = []
    current_section = "Document body"
    buffer: list[str] = []

    def flush() -> None:
        if not buffer:
            return
        joined = " ".join(buffer)
        for chunk in chunk_text(joined):
            out.append((chunk, SourceLocation(section=current_section), {}))
        buffer.clear()

    for para in document.paragraphs:
        text = para.text.strip()
        if not text:
            continue
        if para.style is not None and para.style.name.lower().startswith("heading"):
            flush()
            current_section = text
            continue
        buffer.append(text)
    flush()

    for table_index, table in enumerate(document.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if not rows:
            continue
        table_text = f"Table {table_index}: " + "; ".join(rows)
        for chunk in chunk_text(table_text):
            out.append((
                chunk,
                SourceLocation(section=f"Table {table_index}"),
                {"is_table": True},
            ))
    return out


def _convert_doc_to_docx(path: Path) -> Path:
    """LibreOffice headless conversion. Raises if LibreOffice is absent --
    we never pretend .doc was read."""
    soffice = settings.soffice_cmd
    if not Path(soffice).exists():
        found = shutil.which("soffice")
        if not found:
            raise RuntimeError(
                f"{ErrorCode.CONVERSION_UNAVAILABLE.value}: LibreOffice not found. "
                "Install it to enable legacy .doc support."
            )
        soffice = found

    tmpdir = Path(tempfile.mkdtemp(prefix="evidenceai_doc_"))
    # An isolated user profile lets conversion work even while the user has
    # LibreOffice open; without it the second instance refuses to start.
    # stdin is closed because soffice waits on console input when attached.
    profile = (tmpdir / "profile").as_uri()
    proc = subprocess.run(
        [soffice, "--headless", "--norestore", "--invisible",
         f"-env:UserInstallation={profile}",
         "--convert-to", "docx", "--outdir", str(tmpdir), str(path)],
        capture_output=True,
        text=True,
        timeout=180,
        stdin=subprocess.DEVNULL,
    )
    converted = tmpdir / (path.stem + ".docx")
    if not converted.exists():
        raise RuntimeError(
            f"{ErrorCode.CONVERSION_UNAVAILABLE.value}: LibreOffice conversion "
            f"produced no output. stderr={proc.stderr[:300]}"
        )
    return converted


def _extract_txt(path: Path, warnings: list[str]) -> list[Extracted]:
    text = path.read_text(encoding="utf-8", errors="replace")
    return [(chunk, SourceLocation(), {}) for chunk in chunk_text(text)]


def _extract_csv(path: Path, warnings: list[str]) -> list[Extracted]:
    """Each row becomes one searchable 'column: value' sentence, which
    embeds far better than a raw comma-separated line."""
    import pandas as pd

    frame = pd.read_csv(path)
    columns = [str(c) for c in frame.columns]
    out: list[Extracted] = []
    for row_index, row in frame.iterrows():
        parts = [
            f"{col}: {row[col]}"
            for col in columns
            if str(row[col]).strip() not in ("", "nan", "None")
        ]
        if not parts:
            continue
        text = f"Row {int(row_index) + 1} -- " + "; ".join(parts)
        out.append((
            text,
            SourceLocation(section=f"Row {int(row_index) + 1}"),
            {"row_index": int(row_index)},
        ))
    return out


EXTRACTORS = {
    SourceType.PDF: _extract_pdf,
    SourceType.DOCX: _extract_docx,
    SourceType.TXT: _extract_txt,
    SourceType.CSV: _extract_csv,
}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------

def process_document(path_str: str) -> IngestionResult:
    """Turn one document into an IngestionResult. Never raises: failures come
    back as success=False with an error_code the caller can act on."""
    source: Source | None = None
    warnings: list[str] = []
    try:
        path, source_type = _validate(path_str)
        source = _build_source(path, source_type)

        working_path, working_type = path, source_type
        if source_type is SourceType.DOC:
            working_path = _convert_doc_to_docx(path)
            working_type = SourceType.DOCX
            warnings.append("Converted .doc to .docx via LibreOffice headless.")

        extractor = EXTRACTORS.get(working_type)
        if extractor is None:
            raise ValueError(f"{ErrorCode.UNSUPPORTED_FILE_TYPE.value}: {working_type}")

        extracted = extractor(working_path, warnings)
        if not extracted:
            return IngestionResult(
                source=source,
                items=[],
                success=False,
                error="No extractable text found in document.",
                error_code=ErrorCode.CORRUPT_FILE,
                warnings=warnings,
            )

        items = [
            ContentItem(
                item_id=derive_id(source.source_id, "doc", index),
                source_id=source.source_id,
                modality=Modality.TEXT,
                content=text,
                location=location,
                metadata={
                    "chunk_index": index,
                    "source_type": source_type.value,
                    **extra,
                },
            )
            for index, (text, location, extra) in enumerate(extracted)
        ]
        return IngestionResult(
            source=source,
            items=items,
            relationships=build_document_relationships(items),
            success=True,
            warnings=warnings,
        )

    except Exception as exc:
        logger.error("process_document failed for %s: %s", path_str, exc)
        return _failure(path_str, source, exc, warnings)


def build_document_relationships(items: list[ContentItem]) -> list[Relationship]:
    """SAME_PAGE between chunks sharing a page. SAME_SOURCE is derivable from
    source_id alone and is computed at query time rather than stored O(n^2)."""
    by_page: dict[int, list[ContentItem]] = {}
    for item in items:
        if item.location.page_number is not None:
            by_page.setdefault(item.location.page_number, []).append(item)

    relationships: list[Relationship] = []
    for page, page_items in by_page.items():
        for i in range(len(page_items) - 1):
            a, b = page_items[i], page_items[i + 1]
            relationships.append(Relationship(
                relationship_id=derive_id(a.item_id, b.item_id, "same_page"),
                source_item_id=a.item_id,
                target_item_id=b.item_id,
                relationship_type=RelationshipType.SAME_PAGE,
                confidence=1.0,
                metadata={"page_number": page},
            ))
    return relationships


def _failure(
    path_str: str,
    source: Source | None,
    exc: Exception,
    warnings: list[str],
) -> IngestionResult:
    message = str(exc)
    code = ErrorCode.CORRUPT_FILE
    for candidate in ErrorCode:
        if message.startswith(candidate.value):
            code = candidate
            break

    if source is None:
        path = Path(path_str)
        source = Source(
            source_id=derive_id(path_str),
            filename=path.name or "unknown",
            source_type=resolve_source_type(path),
            file_path=str(path),
        )
    return IngestionResult(
        source=source,
        items=[],
        success=False,
        error=message,
        error_code=code,
        warnings=warnings,
    )
